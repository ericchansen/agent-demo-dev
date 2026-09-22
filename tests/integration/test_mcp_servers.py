"""Raw-wire and SDK-client integration tests for all four local MCP servers.

These tests start each MCP server as a subprocess, communicate via stdio
using the MCP JSON-RPC protocol, and verify that the servers correctly
advertise their tools and handle tool invocations.

Run with:  pytest tests/integration/test_mcp_servers.py -m integration
Skip in fast CI by excluding the integration marker.
"""

from __future__ import annotations

import asyncio
import json
import sys
from collections import deque
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from copy import deepcopy
from itertools import count
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from mcp import Client
from mcp.client.stdio import StdioServerParameters, get_default_environment, stdio_client
from mcp.types import TextContent
from mcp.types.version import LATEST_HANDSHAKE_VERSION
from openpyxl import load_workbook
from pptx import Presentation

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_INITIALIZE_REQUEST: dict[str, Any] = {
    "jsonrpc": "2.0",
    "id": 1,
    "method": "initialize",
    "params": {
        "protocolVersion": "2024-11-05",
        "capabilities": {},
        "clientInfo": {"name": "test-client", "version": "0.1.0"},
    },
}

_INITIALIZED_NOTIFICATION: dict[str, Any] = {
    "jsonrpc": "2.0",
    "method": "notifications/initialized",
}

_LIST_TOOLS_REQUEST: dict[str, Any] = {
    "jsonrpc": "2.0",
    "id": 2,
    "method": "tools/list",
    "params": {},
}


def _call_tool_request(tool_name: str, arguments: dict[str, Any], req_id: int = 3) -> dict[str, Any]:
    return {
        "jsonrpc": "2.0",
        "id": req_id,
        "method": "tools/call",
        "params": {"name": tool_name, "arguments": arguments},
    }


def _encode_message(msg: dict[str, Any]) -> bytes:
    """Encode a JSON-RPC message as newline-delimited JSON for MCP stdio transport."""
    return (json.dumps(msg) + "\n").encode()


async def _read_response(stdout: asyncio.StreamReader) -> dict[str, Any]:
    """Read a single JSON-RPC response line from the MCP server's stdout.

    The MCP Python SDK stdio transport uses newline-delimited JSON:
    one JSON-RPC message per line.
    """
    while True:
        line = await asyncio.wait_for(stdout.readline(), timeout=15)
        if not line:
            raise RuntimeError("MCP server closed stdout unexpectedly")
        decoded = line.decode().strip()
        if not decoded:
            continue
        return json.loads(decoded)


def _server_environment(env_overrides: dict[str, str] | None = None) -> dict[str, str]:
    env = get_default_environment()
    env.update(
        SEARCH_PROVIDER="mock",
        SHAREPOINT_MODE="mock",
        PYTHONPATH=str(Path(__file__).resolve().parents[2]),
        PYTHONDONTWRITEBYTECODE="1",
    )
    env.update(env_overrides or {})
    return env


async def _start_mcp_server(
    module: str, env_overrides: dict[str, str] | None = None, *, cwd: Path | None = None
) -> asyncio.subprocess.Process:
    """Start an MCP server as a subprocess."""
    proc = await asyncio.create_subprocess_exec(
        sys.executable,
        "-m",
        module,
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
        env=_server_environment(env_overrides),
        cwd=cwd,
    )
    return proc


async def _initialize_server(proc: asyncio.subprocess.Process, protocol_version: str = "2024-11-05") -> dict[str, Any]:
    """Send initialize + initialized notification, return the init response."""
    assert proc.stdin is not None
    assert proc.stdout is not None

    request = deepcopy(_INITIALIZE_REQUEST)
    request["params"]["protocolVersion"] = protocol_version
    proc.stdin.write(_encode_message(request))
    await proc.stdin.drain()

    response = await _read_response(proc.stdout)

    # Send the initialized notification (no response expected)
    proc.stdin.write(_encode_message(_INITIALIZED_NOTIFICATION))
    await proc.stdin.drain()

    return response


async def _cleanup(proc: asyncio.subprocess.Process) -> None:
    """Terminate the server process cleanly."""
    if proc.returncode is None:
        proc.terminate()
        try:
            await asyncio.wait_for(proc.wait(), timeout=5)
        except TimeoutError:
            proc.kill()
            await asyncio.wait_for(proc.wait(), timeout=5)


@asynccontextmanager
async def _running_server(module: str, cwd: Path) -> AsyncIterator[asyncio.subprocess.Process]:
    proc = await _start_mcp_server(module, cwd=cwd)
    assert proc.stderr is not None
    stderr_tail: deque[str] = deque(maxlen=30)

    async def drain_stderr() -> None:
        assert proc.stderr is not None
        while line := await proc.stderr.readline():
            stderr_tail.append(line.decode(errors="replace"))

    stderr_task = asyncio.create_task(drain_stderr())
    try:
        yield proc
    finally:
        await _cleanup(proc)
        await stderr_task
        if stderr_tail:
            print("".join(stderr_tail), file=sys.stderr)


async def _exchange(proc: asyncio.subprocess.Process, request: dict[str, Any]) -> dict[str, Any]:
    assert proc.stdin is not None and proc.stdout is not None
    proc.stdin.write(_encode_message(request))
    await proc.stdin.drain()
    response = await _read_response(proc.stdout)
    assert response["jsonrpc"] == "2.0"
    assert response["id"] == request["id"], response
    return response


_SERVER_TOOLS: dict[str, dict[str, dict[str, Any]]] = {
    "researcher": {"research_company": {"company_name": "Tailspin Toys"}},
    "sharepoint": {
        "search_documents": {"query": "Tailspin"},
        "get_document_content": {"drive_id": "sample", "item_id": "sample"},
    },
    "report_generator": {"generate_report": {"title": "Wire report", "customer_name": "Tailspin Toys"}},
    "quota_estimator": {
        "generate_quota_estimation_report": {
            "customer_name": "Tailspin Toys",
            "sales_rows": [
                {"territory": "Northwest", "order_date": "2025-11-01", "revenue": 75000},
                {"territory": "Northwest", "order_date": "2026-05-01", "revenue": 100000},
            ],
        },
    },
}


def _tool_arguments(server_name: str, directory: Path) -> dict[str, dict[str, Any]]:
    tools = deepcopy(_SERVER_TOOLS[server_name])
    if server_name == "quota_estimator":
        tools["generate_quota_estimation_report"]["output_dir"] = str(directory / "quota-artifacts")
    return tools


def _assert_tool_error(response: dict[str, Any], message: str) -> None:
    assert response["result"] == {
        "content": [{"type": "text", "text": message}],
        "isError": True,
    }


def _assert_report_contents(result: dict[str, Any], directory: Path) -> None:
    path = directory / result["file_path"]
    assert path.is_file()
    if result["format"] == "docx":
        text = "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    else:
        text = "\n".join(
            shape.text for slide in Presentation(str(path)).slides for shape in slide.shapes if shape.has_text_frame
        )
    assert "Tailspin Toys" in text
    assert "Sources & Citations" in text
    assert "Fabric pipeline query" in text


def _assert_quota_contents(result: dict[str, Any]) -> None:
    assert result["status"] == "success"
    assert set(result["artifacts"]) == {"xlsx", "html", "pdf"}
    assert any("SalesOrderHeader" in citation for citation in result["citations"])
    workbook = load_workbook(result["artifacts"]["xlsx"], read_only=True, data_only=True)
    try:
        assert workbook["Summary"]["A1"].value == "Quota Estimate - Tailspin Toys"
        assert sum(row[3] for row in workbook["Sales Detail"].iter_rows(min_row=2, values_only=True)) == 175000
        assert "SalesOrderHeader" in workbook["Methodology"]["A8"].value
    finally:
        workbook.close()
    html = Path(result["artifacts"]["html"]).read_text(encoding="utf-8")
    assert "Tailspin Toys" in html and "Northwest" in html
    assert "Citations" in html and "SalesOrderHeader" in html
    pdf = Path(result["artifacts"]["pdf"]).read_bytes()
    assert pdf.startswith(b"%PDF-") and pdf.rstrip().endswith(b"%%EOF")


@pytest.mark.integration
@pytest.mark.asyncio
@pytest.mark.parametrize("server_name", list(_SERVER_TOOLS))
@pytest.mark.parametrize("protocol_version", ["2024-11-05", LATEST_HANDSHAKE_VERSION])
async def test_wire_validation_and_recovery(server_name: str, protocol_version: str, tmp_path: Path) -> None:
    module = f"src.agents.{server_name}.mcp_server"
    tools = _tool_arguments(server_name, tmp_path)
    request_ids = count(2)
    async with _running_server(module, tmp_path) as proc:
        initialized = await _initialize_server(proc, protocol_version)
        assert initialized["result"]["protocolVersion"] == protocol_version

        unknown = await _exchange(proc, _call_tool_request("missing_tool", {}, next(request_ids)))
        _assert_tool_error(unknown, "Unknown tool: missing_tool")
        listed = await _exchange(
            proc, {"jsonrpc": "2.0", "id": next(request_ids), "method": "tools/list", "params": {}}
        )
        schemas = {tool["name"]: tool["inputSchema"] for tool in listed["result"]["tools"]}
        assert set(schemas) == set(tools)
        assert all("input_schema" not in tool for tool in listed["result"]["tools"])

        for name, arguments in tools.items():
            required = list(_SERVER_TOOLS[server_name][name])
            assert schemas[name]["required"] == required
            first = required[0]
            for empty_params in ({"name": name}, {"name": name, "arguments": None}, {"name": name, "arguments": {}}):
                response = await _exchange(
                    proc, {"jsonrpc": "2.0", "id": next(request_ids), "method": "tools/call", "params": empty_params}
                )
                _assert_tool_error(response, f"Input validation error: '{first}' is a required property")
            for missing in required:
                invalid = {key: value for key, value in arguments.items() if key != missing}
                response = await _exchange(proc, _call_tool_request(name, invalid, next(request_ids)))
                _assert_tool_error(response, f"Input validation error: '{missing}' is a required property")
            for value in (42, None):
                response = await _exchange(
                    proc, _call_tool_request(name, {**arguments, first: value}, next(request_ids))
                )
                _assert_tool_error(response, f"Input validation error: {value} is not of type 'string'")
            for invalid_params in (
                {"name": name, "arguments": []},
                {"name": name, "arguments": ["bad"]},
                {"name": name, "arguments": "bad"},
                {"name": 42, "arguments": arguments},
                {"arguments": arguments},
            ):
                response = await _exchange(
                    proc, {"jsonrpc": "2.0", "id": next(request_ids), "method": "tools/call", "params": invalid_params}
                )
                assert response["error"] == {"code": -32602, "message": "Invalid request parameters", "data": ""}
            if server_name in {"report_generator", "quota_estimator"}:
                response = await _exchange(
                    proc, _call_tool_request(name, {**arguments, "unexpected": True}, next(request_ids))
                )
                _assert_tool_error(
                    response,
                    "Input validation error: Additional properties are not allowed ('unexpected' was unexpected)",
                )

        assert not (tmp_path / "output").exists()
        assert not (tmp_path / "quota-artifacts").exists()
        unknown_method = await _exchange(
            proc, {"jsonrpc": "2.0", "id": next(request_ids), "method": "missing/method", "params": {}}
        )
        assert unknown_method["error"] == {"code": -32601, "message": "Method not found", "data": "missing/method"}

        # Malformed outer envelopes are discarded by the transport, not correlated tool errors.
        assert proc.stdin is not None and proc.stdout is not None
        rejected_id = next(request_ids)
        proc.stdin.write(b"{broken json\n")
        proc.stdin.write(_encode_message({"jsonrpc": "2.0", "id": rejected_id, "method": "tools/call", "params": []}))
        ping_id = next(request_ids)
        ping = await _exchange(proc, {"jsonrpc": "2.0", "id": ping_id, "method": "ping", "params": {}})
        assert ping == {"jsonrpc": "2.0", "id": ping_id, "result": {}}

        for name, arguments in tools.items():
            if server_name in {"researcher", "sharepoint"}:
                arguments = {**arguments, "unexpected": True}
            response = await _exchange(proc, _call_tool_request(name, arguments, next(request_ids)))
            assert response["result"].get("isError", False) is False
            assert "resultType" not in response["result"]
            data = json.loads(response["result"]["content"][0]["text"])
            if server_name == "report_generator":
                _assert_report_contents(data, tmp_path)
                response = await _exchange(
                    proc, _call_tool_request(name, {**arguments, "format": "pptx"}, next(request_ids))
                )
                assert response["result"].get("isError", False) is False
                _assert_report_contents(json.loads(response["result"]["content"][0]["text"]), tmp_path)
            elif server_name == "quota_estimator":
                _assert_quota_contents(data)
            elif name == "research_company":
                assert data["company_name"] == "Tailspin Toys" and data["articles"]
            elif name == "search_documents":
                assert data and "Tailspin" in data[0]["name"]
            else:
                assert "Tailspin Toys" in data["content_text"]


@pytest.mark.integration
@pytest.mark.asyncio
@pytest.mark.parametrize("server_name", list(_SERVER_TOOLS))
async def test_real_mcp2_client_interoperability(server_name: str, tmp_path: Path) -> None:
    tools = _tool_arguments(server_name, tmp_path)
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", f"src.agents.{server_name}.mcp_server"],
        env=_server_environment(),
        cwd=tmp_path,
    )
    with (tmp_path / "server-stderr.log").open("w", encoding="utf-8") as stderr:
        async with asyncio.timeout(15), Client(stdio_client(params, errlog=stderr)) as client:
            discovered = await client.list_tools()
            assert {tool.name for tool in discovered.tools} == set(tools)
            assert all(tool.input_schema["type"] == "object" for tool in discovered.tools)
            rejected = await client.call_tool("missing_tool", {})
            assert rejected.is_error is True
            for name, arguments in tools.items():
                response = await client.call_tool(name, arguments)
                assert response.is_error is False
                assert isinstance(response.content[0], TextContent)
                assert json.loads(response.content[0].text)


# ---------------------------------------------------------------------------
# Researcher Agent Tests
# ---------------------------------------------------------------------------


@pytest.mark.integration
@pytest.mark.asyncio
async def test_researcher_list_tools():
    """Researcher MCP server advertises the research_company tool."""
    proc = await _start_mcp_server(
        "src.agents.researcher.mcp_server",
        env_overrides={"SEARCH_PROVIDER": "mock"},
    )
    try:
        await _initialize_server(proc)

        assert proc.stdin is not None
        assert proc.stdout is not None

        proc.stdin.write(_encode_message(_LIST_TOOLS_REQUEST))
        await proc.stdin.drain()

        response = await _read_response(proc.stdout)

        assert "result" in response, f"Expected 'result' in response, got: {response}"
        tools = response["result"]["tools"]
        assert len(tools) >= 1

        tool_names = [t["name"] for t in tools]
        assert "research_company" in tool_names

        # Verify schema shape
        research_tool = next(t for t in tools if t["name"] == "research_company")
        assert "inputSchema" in research_tool
        assert "company_name" in research_tool["inputSchema"]["properties"]
    finally:
        await _cleanup(proc)


@pytest.mark.integration
@pytest.mark.asyncio
async def test_researcher_call_tool_mock():
    """Researcher MCP server returns mock data for Tailspin Toys."""
    proc = await _start_mcp_server(
        "src.agents.researcher.mcp_server",
        env_overrides={"SEARCH_PROVIDER": "mock"},
    )
    try:
        await _initialize_server(proc)

        assert proc.stdin is not None
        assert proc.stdout is not None

        request = _call_tool_request("research_company", {"company_name": "Tailspin Toys"})
        proc.stdin.write(_encode_message(request))
        await proc.stdin.drain()

        response = await _read_response(proc.stdout)

        assert "result" in response, f"Expected 'result' in response, got: {response}"
        content = response["result"]["content"]
        assert len(content) >= 1
        assert content[0]["type"] == "text"

        data = json.loads(content[0]["text"])
        assert data["company_name"] == "Tailspin Toys"
        assert len(data["articles"]) > 0
        assert "key_metrics" in data
    finally:
        await _cleanup(proc)


@pytest.mark.integration
@pytest.mark.asyncio
async def test_researcher_call_tool_unknown_company():
    """Researcher MCP server returns a generic response for an unknown company in mock mode."""
    proc = await _start_mcp_server(
        "src.agents.researcher.mcp_server",
        env_overrides={"SEARCH_PROVIDER": "mock"},
    )
    try:
        await _initialize_server(proc)

        assert proc.stdin is not None
        assert proc.stdout is not None

        request = _call_tool_request("research_company", {"company_name": "UnknownCorp"})
        proc.stdin.write(_encode_message(request))
        await proc.stdin.drain()

        response = await _read_response(proc.stdout)

        assert "result" in response
        data = json.loads(response["result"]["content"][0]["text"])
        assert data["company_name"] == "UnknownCorp"
        assert data["articles"] == []
    finally:
        await _cleanup(proc)


# ---------------------------------------------------------------------------
# SharePoint Agent Tests
# ---------------------------------------------------------------------------


@pytest.mark.integration
@pytest.mark.asyncio
async def test_sharepoint_list_tools():
    """SharePoint MCP server advertises search_documents and get_document_content tools."""
    proc = await _start_mcp_server(
        "src.agents.sharepoint.mcp_server",
        env_overrides={"SHAREPOINT_MODE": "mock"},
    )
    try:
        await _initialize_server(proc)

        assert proc.stdin is not None
        assert proc.stdout is not None

        proc.stdin.write(_encode_message(_LIST_TOOLS_REQUEST))
        await proc.stdin.drain()

        response = await _read_response(proc.stdout)

        assert "result" in response, f"Expected 'result' in response, got: {response}"
        tools = response["result"]["tools"]
        assert len(tools) >= 2

        tool_names = [t["name"] for t in tools]
        assert "search_documents" in tool_names
        assert "get_document_content" in tool_names

        # Verify schema for search_documents
        search_tool = next(t for t in tools if t["name"] == "search_documents")
        assert "query" in search_tool["inputSchema"]["properties"]
    finally:
        await _cleanup(proc)


@pytest.mark.integration
@pytest.mark.asyncio
async def test_sharepoint_search_documents_mock():
    """SharePoint MCP server returns mock documents for a matching query."""
    proc = await _start_mcp_server(
        "src.agents.sharepoint.mcp_server",
        env_overrides={"SHAREPOINT_MODE": "mock"},
    )
    try:
        await _initialize_server(proc)

        assert proc.stdin is not None
        assert proc.stdout is not None

        request = _call_tool_request("search_documents", {"query": "Tailspin"})
        proc.stdin.write(_encode_message(request))
        await proc.stdin.drain()

        response = await _read_response(proc.stdout)

        assert "result" in response, f"Expected 'result' in response, got: {response}"
        content = response["result"]["content"]
        assert len(content) >= 1
        assert content[0]["type"] == "text"

        data = json.loads(content[0]["text"])
        assert isinstance(data, list)
        assert len(data) > 0
        assert "name" in data[0]
        assert "Tailspin" in data[0]["name"]
    finally:
        await _cleanup(proc)


@pytest.mark.integration
@pytest.mark.asyncio
async def test_sharepoint_search_documents_no_results():
    """SharePoint MCP server returns an empty list for a non-matching query."""
    proc = await _start_mcp_server(
        "src.agents.sharepoint.mcp_server",
        env_overrides={"SHAREPOINT_MODE": "mock"},
    )
    try:
        await _initialize_server(proc)

        assert proc.stdin is not None
        assert proc.stdout is not None

        request = _call_tool_request("search_documents", {"query": "nonexistent-xyz-12345"})
        proc.stdin.write(_encode_message(request))
        await proc.stdin.drain()

        response = await _read_response(proc.stdout)

        assert "result" in response
        data = json.loads(response["result"]["content"][0]["text"])
        assert isinstance(data, list)
        assert len(data) == 0
    finally:
        await _cleanup(proc)


# ---------------------------------------------------------------------------
# Quota Estimator Tests
# ---------------------------------------------------------------------------


@pytest.mark.integration
@pytest.mark.asyncio
async def test_quota_estimator_list_tools():
    """Quota estimator MCP server advertises the quota artifact generation tool."""
    proc = await _start_mcp_server("src.agents.quota_estimator.mcp_server")
    try:
        await _initialize_server(proc)

        assert proc.stdin is not None
        assert proc.stdout is not None

        proc.stdin.write(_encode_message(_LIST_TOOLS_REQUEST))
        await proc.stdin.drain()

        response = await _read_response(proc.stdout)

        assert "result" in response, f"Expected 'result' in response, got: {response}"
        tool_names = [tool["name"] for tool in response["result"]["tools"]]
        assert "generate_quota_estimation_report" in tool_names
    finally:
        await _cleanup(proc)


@pytest.mark.integration
@pytest.mark.asyncio
async def test_quota_estimator_call_tool(tmp_path):
    """Quota estimator MCP server generates XLSX, HTML, and PDF artifacts over stdio."""
    proc = await _start_mcp_server("src.agents.quota_estimator.mcp_server")
    try:
        await _initialize_server(proc)

        assert proc.stdin is not None
        assert proc.stdout is not None

        request = _call_tool_request(
            "generate_quota_estimation_report",
            {
                "customer_name": "Tailspin Toys",
                "sales_rows": [
                    {
                        "territory": "Northwest",
                        "category": "Novelty Items",
                        "order_date": "2025-11-01",
                        "revenue": 75000,
                        "quantity": 180,
                    },
                    {
                        "territory": "Northwest",
                        "category": "Novelty Items",
                        "order_date": "2026-05-01",
                        "revenue": 100000,
                        "quantity": 250,
                    },
                ],
                "research_data": {"summary": "Retail demand is expanding 10%."},
                "workiq_activity": {"engagement_score": "High", "recent_activity": [{"type": "meeting"}]},
                "scenario": "aggressive",
                "output_dir": str(tmp_path),
            },
        )
        proc.stdin.write(_encode_message(request))
        await proc.stdin.drain()

        response = await _read_response(proc.stdout)

        assert "result" in response, f"Expected 'result' in response, got: {response}"
        data = json.loads(response["result"]["content"][0]["text"])
        assert data["status"] == "success"
        assert data["scenario"] == "aggressive"
        assert set(data["artifacts"]) == {"xlsx", "html", "pdf"}
    finally:
        await _cleanup(proc)
